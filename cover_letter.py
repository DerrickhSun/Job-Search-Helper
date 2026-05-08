"""
Cover Letter Generator
Uses a DSPy ChainOfThought module to write a tailored 3-paragraph cover letter for each job.

Optional ``example_cover_letter`` on the resume dict: a cover letter the candidate has already used or
sent before. The model may reuse its wording, role descriptions, and employers when that field is set
(see ``COVER_INSTRUCTION``).
"""

from __future__ import annotations

import logging
import re
from pathlib import Path
from typing import Any

import dspy

log = logging.getLogger(__name__)

_EXPERIENCE_BLOCK_MAX_CHARS = 20_000
_EXAMPLE_COVER_LETTER_MAX_CHARS = 8_000
_MAX_BULLETS_PER_ROLE = 80


def normalize_cover_letter_dashes(text: str) -> str:
    """Replace em/en dashes and similar with ASCII hyphen-minus (forms often reject fancy punctuation)."""
    if not text:
        return text
    t = text.replace("\u2014", "-").replace("\u2013", "-").replace("\u2015", "-")
    t = t.replace("\u2212", "-")  # minus sign
    return t


def write_cover_letter_docx(body: str, path: Path | str) -> Path:
    """
    Write cover letter body to a ``.docx`` file (paragraphs split on blank lines).
    Requires ``python-docx``.
    """
    from docx import Document

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    body = normalize_cover_letter_dashes(body.strip())
    doc = Document()
    parts = [p.strip() for p in re.split(r"\n\s*\n", body) if p.strip()]
    if not parts:
        doc.add_paragraph(body or "")
    else:
        for p in parts:
            doc.add_paragraph(p)
    doc.save(str(path.resolve()))
    return path


COVER_INSTRUCTION = """Write a concise, professional cover letter for this job application.

Requirements:
- Exactly 3 short paragraphs, no fluff
- Paragraph 1: Why this role fits and your relevant background (name this employer and role using
  ``job_title`` and ``job_company`` where natural).
- Paragraph 2: One specific example of relevant past work or impact. How you source it depends on
  ``example_cover_letter``:
  * If ``example_cover_letter`` is **non-empty**, treat it as the candidate's own prior cover letter.
    You may **reuse, adapt, or copy** its sentences, role descriptions, employers, products, and
    accomplishments when they still fit this application. It is trustworthy material from the user.
  * If ``example_cover_letter`` is **empty**, ground this paragraph only in ``summary``, ``skills``, and
    ``experience`` below — do not invent employers, titles, projects, technologies, or metrics that
    are not clearly supported there.
- Paragraph 3: Enthusiasm for **this** company and role and a clear call to action (use ``job_title``,
  ``job_company``, and ``job_description`` so the letter targets the current posting).
- Tone: confident but not arrogant, conversational but professional
- Do NOT include a date, address block, or "Dear Hiring Manager" — only the three body paragraphs
- Keep under 220 words total

If ``example_cover_letter`` is empty, ignore the non-empty branch above for paragraph 2."""


class CoverLetterModule(dspy.Module):
    def __init__(self):
        super().__init__()
        self.write = dspy.ChainOfThought(
            "instruction, name, email, summary, skills, experience, example_cover_letter, "
            "job_title, job_company, job_description -> cover_letter: str"
        )

    def forward(
        self,
        name: str,
        email: str,
        summary: str,
        skills: str,
        experience: str,
        example_cover_letter: str,
        job_title: str,
        job_company: str,
        job_description: str,
    ):
        return self.write(
            instruction=COVER_INSTRUCTION,
            name=name,
            email=email,
            summary=summary,
            skills=skills,
            experience=experience,
            example_cover_letter=example_cover_letter,
            job_title=job_title,
            job_company=job_company,
            job_description=job_description,
        )


def _format_experience_for_cover_letter(resume: dict[str, Any], *, max_chars: int = _EXPERIENCE_BLOCK_MAX_CHARS) -> str:
    """Serialize every experience entry (title, company, dates, bullets) for the model."""
    lines: list[str] = []
    for r in resume.get("experience") or []:
        if not isinstance(r, dict):
            continue
        title = str(r.get("title") or "").strip()
        company = str(r.get("company") or "").strip()
        dates = str(r.get("dates") or "").strip()
        head_bits = [x for x in (title, company, dates) if x]
        if head_bits:
            lines.append(" | ".join(head_bits))
        bullets = r.get("bullets") or []
        if isinstance(bullets, (list, tuple)):
            for b in bullets[:_MAX_BULLETS_PER_ROLE]:
                bt = str(b).strip()
                if bt:
                    lines.append(f"  - {bt}")
        lines.append("")
    text = "\n".join(lines).strip()
    if len(text) > max_chars:
        text = text[: max_chars - 48].rstrip() + "\n[... experience truncated for model input length ...]"
    return text or "Not provided"


class CoverLetterGenerator:
    def __init__(self):
        self._module = CoverLetterModule()

    def generate(self, resume: dict, job: dict) -> str:
        """
        Generates and returns a cover letter string.
        Falls back to a template if the DSPy call fails.
        """
        try:
            return self._generate_with_dspy(resume, job)
        except Exception as e:
            log.warning("Cover letter generation failed (%s), using template fallback", e)
            return self._template_fallback(resume, job)

    def _generate_with_dspy(self, resume: dict, job: dict) -> str:
        example = (resume.get("example_cover_letter") or "").strip()
        if len(example) > _EXAMPLE_COVER_LETTER_MAX_CHARS:
            example = example[: _EXAMPLE_COVER_LETTER_MAX_CHARS - 40].rstrip() + "\n[... truncated ...]"

        result = self._module(
            name=resume.get("name", ""),
            email=resume.get("email", ""),
            summary=(resume.get("summary", "") or "")[:500],
            skills=", ".join(resume.get("skills", [])[:12]),
            experience=_format_experience_for_cover_letter(resume),
            example_cover_letter=example,
            job_title=job.get("title", ""),
            job_company=job.get("company", ""),
            job_description=(job.get("description", ""))[:4000],
        )

        return normalize_cover_letter_dashes(result.cover_letter.strip())

    def _template_fallback(self, resume: dict, job: dict) -> str:
        name = resume.get("name", "Candidate")
        skills = ", ".join(resume.get("skills", [])[:5])
        return normalize_cover_letter_dashes(
            f"I am excited to apply for the {job.get('title')} position at {job.get('company')}. "
            f"With expertise in {skills}, I believe I would be a strong addition to your team.\n\n"
            f"Throughout my career, I have consistently delivered results in similar roles and am confident "
            f"I can bring that same dedication to {job.get('company')}.\n\n"
            f"I would welcome the opportunity to discuss how my background aligns with your needs. "
            f"Thank you for considering my application.\n\nSincerely,\n{name}"
        )
