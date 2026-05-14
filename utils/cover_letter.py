"""
Cover Letter Generator
Uses a DSPy ChainOfThought module to write a tailored four-paragraph cover letter for each job.

Optional ``example_cover_letter`` on the resume dict: a cover letter the candidate has already used or
sent before. The model may reuse its tone, structure, and wording when that field is set
(see ``COVER_INSTRUCTION``). Optional ``projects``: list of ``{title, dates, description}`` (same shape
as produced by :class:`utils.resume_parser.ResumeParser`) is passed through to the model for grounding.
"""

from __future__ import annotations

import logging
import re
from pathlib import Path
from typing import Any

import dspy

from .resume_parser import experience_entry_description, project_entry_description

log = logging.getLogger(__name__)

_EXPERIENCE_BLOCK_MAX_CHARS = 20_000
_PROJECTS_BLOCK_MAX_CHARS = 12_000
_EXAMPLE_COVER_LETTER_MAX_CHARS = 8_000
_MAX_EXPERIENCE_DESCRIPTION_CHARS_PER_ROLE = 8000
_MAX_PROJECT_DESCRIPTION_CHARS = 6000


def normalize_cover_letter_dashes(text: str) -> str:
    """Replace em/en dashes and similar with ASCII hyphen-minus (forms often reject fancy punctuation)."""
    if not text:
        return text
    t = text.replace("\u2014", "-").replace("\u2013", "-").replace("\u2015", "-")
    t = t.replace("\u2212", "-")  # minus sign
    return t


# First line must introduce the applicant by name, then the apply-for clause.
_COVER_OPENING_OK = re.compile(
    r"My name is\b.+,?\s*and\s+(?:I am writing to apply for|I write to apply for)\b",
    re.IGNORECASE,
)

# Legacy apply-only opening (no "My name is") — removed when prepending the canonical intro.
_LEGACY_APPLY_START = re.compile(
    r"^\s*(?:"
    r"I am writing to apply for\b"
    r"|I write to apply for\b"
    r"|I am applying for\b"
    r"|I am writing to express interest in\b"
    r")[^.!?]*[.!?]\s*",
    re.IGNORECASE | re.DOTALL,
)


def _canonical_opening_line(applicant_name: str, job_title: str, job_company: str) -> str:
    name = (applicant_name or "").strip() or "Candidate"
    title = (job_title or "").strip() or "this position"
    company = (job_company or "").strip() or "your organization"
    return f"My name is {name}, and I am writing to apply for the {title} position at {company}."


def _ensure_apply_opening_sentence(
    text: str,
    applicant_name: str,
    job_title: str,
    job_company: str,
) -> str:
    """
    Ensure the letter begins with: My name is {name}, and I am writing to apply for the {title} at {company}.
    If the model omitted the name intro or used a legacy-only opening, fix it.
    """
    t = normalize_cover_letter_dashes((text or "").strip())
    canonical = _canonical_opening_line(applicant_name, job_title, job_company)
    if not t:
        return canonical

    first_line = t.split("\n", 1)[0]
    if _COVER_OPENING_OK.search(first_line):
        return t

    m = _LEGACY_APPLY_START.match(t)
    if m:
        rest = t[m.end() :].strip()
        return normalize_cover_letter_dashes(f"{canonical} {rest}".strip()) if rest else canonical

    return normalize_cover_letter_dashes(f"{canonical} {t}".strip())


def write_cover_letter_docx(body: str, path: Path | str) -> Path:
    """
    Write cover letter body to a ``.docx`` file (paragraphs split on blank lines).
    Requires ``python-docx``.
    """
    from docx import Document

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    body = normalize_cover_letter_dashes(body.strip())
    doc = Document()
    parts = [p.strip() for p in re.split(r"\n\s*\n", body) if p.strip()]
    if not parts:
        doc.add_paragraph(body or "")
    else:
        for p in parts:
            doc.add_paragraph(p)
    doc.save(str(path.resolve()))
    return path


_MAX_COVER_LETTER_FILENAME_STEM_CHARS = 200


def _normalize_cover_letter_site(site: str) -> str:
    s = (site or "").strip().lower()
    if s in ("greenhouse", "gh") or "greenhouse" in s:
        return "greenhouse"
    return "linkedin"


def _sanitize_cover_letter_filename_segment(s: str, max_len: int) -> str:
    t = (s or "").strip()
    t = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "", t)
    t = re.sub(r"\s+", "_", t)
    t = re.sub(r"_+", "_", t).strip("._-")
    if not t:
        t = "unknown"
    return t[:max_len].rstrip("._-")


def cover_letter_docx_stem(
    *,
    site: str,
    company: str,
    title: str,
    job_id: str,
) -> str:
    """
    Filesystem-safe filename **stem** (no ``.docx``) for a cover letter:
    ``{site}_{company}_{title}_{job_id}`` (``site`` is ``linkedin`` or ``greenhouse``).
    """
    board = _normalize_cover_letter_site(site)
    co = _sanitize_cover_letter_filename_segment(company or "Company", 55)
    ti = _sanitize_cover_letter_filename_segment(title or "Position", 75)
    raw_id = str(job_id or "job").strip()
    jid = re.sub(r"[^\w\-.]+", "_", raw_id).strip("_")
    jid = (jid or "job")[:48]
    stem = "_".join((board, co, ti, jid))
    stem = re.sub(r"_+", "_", stem)
    if len(stem) > _MAX_COVER_LETTER_FILENAME_STEM_CHARS:
        stem = stem[:_MAX_COVER_LETTER_FILENAME_STEM_CHARS].rstrip("._-")
    return stem


def cover_letter_docx_path_unique(
    output_dir: Path | str,
    *,
    site: str,
    company: str,
    title: str,
    job_id: str,
) -> Path:
    """
    ``output_dir / {stem}.docx`` using :func:`cover_letter_docx_stem`; if that path exists, use
    ``{stem}__2.docx``, ``{stem}__3.docx``, …
    """
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    stem = cover_letter_docx_stem(site=site, company=company, title=title, job_id=str(job_id))
    path = output_dir / f"{stem}.docx"
    if not path.exists():
        return path
    n = 2
    while True:
        cand = output_dir / f"{stem}__{n}.docx"
        if not cand.exists():
            return cand
        n += 1


COVER_INSTRUCTION = """Write a concise, professional cover letter for this job application.

Requirements:
- Exactly **four** paragraphs (plain body text only). No fluff; keep each paragraph focused.
- **Opening sentence (mandatory):** The entire cover letter must **begin** with one clear sentence that is a
  variant of: "My name is ``name``, and I am writing to apply for the ``job_title`` position at ``job_company``."
  Use the applicant's full ``name`` from the inputs, plus the actual ``job_title`` and ``job_company`` strings.
  Acceptable small variants: use "I write to apply for" instead of "I am writing to apply for" after the comma;
  optional comma before "and". Do not put any text before this opening sentence.
- Paragraph 1 — Introduction: After that opening sentence, briefly introduce the applicant's **education**
  (degree, field, stage such as current student or recent graduate when supported by ``summary`` / ``experience``).
  Acknowledge **this** employer's mission, product direction, or stated goals using ``job_company``, ``job_title``,
  and ``job_description`` (what the company is trying to achieve or build — not generic praise).
- Paragraph 2 — Fit through experience: Include **three** distinct experiences or projects that show the
  applicant is a good fit for **this** role. Each should be a sentence or two. How you source them:
  * If ``example_cover_letter`` is **non-empty**, treat it as the user's own sample letter (tone and structure
    guide). You may **reuse, adapt, or echo** its phrasing when it still fits; prefer grounding specifics in
    ``experience`` and ``projects`` below when facts conflict.
  * Always anchor claims in ``summary``, ``skills``, ``experience`` (per-role descriptions), and ``projects``
    — do not invent employers, titles, projects, technologies, or metrics not clearly supported there.
  If fewer than three solid items exist in the materials, use the strongest available items once each and do
  not invent a third.
- Paragraph 3 — Company and role: Focus on **this** company and job. Explain what specifically draws the
  applicant (mission, product, tech stack, team scope from ``job_description``) and why their background is a
  strong match for ``job_title`` at ``job_company``.
- Paragraph 4 — Conclusion: A **brief** closing that thanks the reader for their time (and consideration if
  natural). No long repetition of paragraph 3.
- Tone: confident but not arrogant, conversational but professional.
- Do NOT include a date, address block, salutation (e.g. "Dear Hiring Manager"), or signature line — only the
  four body paragraphs.
- Aim for under 320 words total.

If ``example_cover_letter`` is empty, ignore any branch that refers to reusing its wording except as optional
style inspiration; still write paragraph 2 from ``summary``, ``skills``, ``experience``, and ``projects`` only."""


class CoverLetterModule(dspy.Module):
    def __init__(self):
        super().__init__()
        self.write = dspy.ChainOfThought(
            "instruction, name, email, summary, skills, experience, projects, example_cover_letter, "
            "job_title, job_company, job_description -> cover_letter: str"
        )

    def forward(
        self,
        name: str,
        email: str,
        summary: str,
        skills: str,
        experience: str,
        projects: str,
        example_cover_letter: str,
        job_title: str,
        job_company: str,
        job_description: str,
    ):
        return self.write(
            instruction=COVER_INSTRUCTION,
            name=name,
            email=email,
            summary=summary,
            skills=skills,
            experience=experience,
            projects=projects,
            example_cover_letter=example_cover_letter,
            job_title=job_title,
            job_company=job_company,
            job_description=job_description,
        )


def _format_experience_for_cover_letter(resume: dict[str, Any], *, max_chars: int = _EXPERIENCE_BLOCK_MAX_CHARS) -> str:
    """Serialize every experience entry (title, company, dates, description) for the model."""
    lines: list[str] = []
    for r in resume.get("experience") or []:
        if not isinstance(r, dict):
            continue
        title = str(r.get("title") or "").strip()
        company = str(r.get("company") or "").strip()
        dates = str(r.get("dates") or "").strip()
        head_bits = [x for x in (title, company, dates) if x]
        if head_bits:
            lines.append(" | ".join(head_bits))
        body = experience_entry_description(r)
        if len(body) > _MAX_EXPERIENCE_DESCRIPTION_CHARS_PER_ROLE:
            body = (
                body[: _MAX_EXPERIENCE_DESCRIPTION_CHARS_PER_ROLE - 48].rstrip()
                + "\n[... role description truncated ...]"
            )
        if body:
            lines.append(body)
        lines.append("")
    text = "\n".join(lines).strip()
    if len(text) > max_chars:
        text = text[: max_chars - 48].rstrip() + "\n[... experience truncated for model input length ...]"
    return text or "Not provided"


def _format_projects_for_cover_letter(resume: dict[str, Any], *, max_chars: int = _PROJECTS_BLOCK_MAX_CHARS) -> str:
    """Serialize every project entry (title, dates, description) for the model."""
    lines: list[str] = []
    for p in resume.get("projects") or []:
        if not isinstance(p, dict):
            continue
        title = str(p.get("title") or "").strip()
        dates = str(p.get("dates") or "").strip()
        head_bits = [x for x in (title, dates) if x]
        if head_bits:
            lines.append(" | ".join(head_bits))
        body = project_entry_description(p)
        if len(body) > _MAX_PROJECT_DESCRIPTION_CHARS:
            body = (
                body[: _MAX_PROJECT_DESCRIPTION_CHARS - 48].rstrip()
                + "\n[... project description truncated ...]"
            )
        if body:
            lines.append(body)
        lines.append("")
    text = "\n".join(lines).strip()
    if len(text) > max_chars:
        text = text[: max_chars - 48].rstrip() + "\n[... projects truncated for model input length ...]"
    return text or "Not provided"


class CoverLetterGenerator:
    def __init__(self):
        self._module = CoverLetterModule()

    def generate(self, resume: dict, job: dict) -> str:
        """
        Generates and returns a cover letter string.
        Falls back to a template if the DSPy call fails.
        """
        try:
            return self._generate_with_dspy(resume, job)
        except Exception as e:
            log.warning("Cover letter generation failed (%s), using template fallback", e)
            return self._template_fallback(resume, job)

    def _generate_with_dspy(self, resume: dict, job: dict) -> str:
        example = (resume.get("example_cover_letter") or "").strip()
        if len(example) > _EXAMPLE_COVER_LETTER_MAX_CHARS:
            example = example[: _EXAMPLE_COVER_LETTER_MAX_CHARS - 40].rstrip() + "\n[... truncated ...]"

        result = self._module(
            name=resume.get("name", ""),
            email=resume.get("email", ""),
            summary=(resume.get("summary", "") or "")[:500],
            skills=", ".join(resume.get("skills", [])[:12]),
            experience=_format_experience_for_cover_letter(resume),
            projects=_format_projects_for_cover_letter(resume),
            example_cover_letter=example,
            job_title=job.get("title", ""),
            job_company=job.get("company", ""),
            job_description=(job.get("description", ""))[:4000],
        )

        raw = normalize_cover_letter_dashes(result.cover_letter.strip())
        return _ensure_apply_opening_sentence(
            raw,
            str(resume.get("name") or ""),
            str(job.get("title") or ""),
            str(job.get("company") or ""),
        )

    def _template_fallback(self, resume: dict, job: dict) -> str:
        title = job.get("title") or "this role"
        company = job.get("company") or "your organization"
        skills = ", ".join(resume.get("skills", [])[:5])
        summary = (resume.get("summary") or "").strip()
        edu_hint = summary[:200] + ("..." if len(summary) > 200 else "") if summary else (
            f"My background includes strengths in {skills}."
        )
        display_name = str(resume.get("name") or "").strip() or "Candidate"
        open_line = _canonical_opening_line(display_name, str(title), str(company))
        return _ensure_apply_opening_sentence(
            "\n\n".join(
                [
                    f"{open_line} {edu_hint} "
                    f"I am motivated by opportunities where my training can support teams building impactful products.",
                    f"Relevant experience includes work described in my resume across software engineering, collaboration, "
                    f"and technical depth in areas such as {skills}. "
                    f"I have applied these skills in course projects, internships, and hands-on development work.",
                    f"This role at {company} aligns with my interests and the problems I want to solve; I am eager to "
                    f"contribute to your goals for the {title} position and grow with the team.",
                    f"Thank you for your time and consideration.",
                ]
            ),
            display_name,
            str(job.get("title") or ""),
            str(job.get("company") or ""),
        )
